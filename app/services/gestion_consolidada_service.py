from __future__ import annotations

from pathlib import Path
from datetime import datetime
import os
import re
from typing import Any

# === CONSOLIDAR_V2_DATA_DIR_RUTAS_FIX_BEGIN ===

def _gc_v2_resolver_archivo_aster_gestion(base, fecha):
    """
    ASTER actual:
        E:\data\YYYYMMDD\Aster\Gestion\YYYYMMDD_Gestion_aster.xlsx

    ASTER legado:
        E:\data\YYYYMMDD\Aster\aster_YYYYMMDD\Gestion\YYYYMMDD_Gestion_aster.xlsx
    """
    from pathlib import Path

    base = Path(base)

    actual = base / "Aster" / "Gestion" / f"{fecha}_Gestion_aster.xlsx"
    legado = base / "Aster" / f"aster_{fecha}" / "Gestion" / f"{fecha}_Gestion_aster.xlsx"

    if actual.exists():
        return actual

    if legado.exists():
        return legado

    # Devolver actual como ruta esperada para que la pantalla muestre dónde debe estar.
    return actual


def _gc_v2_resolver_archivo_orion_gestion(base, fecha):
    """
    ORION:
        E:\data\YYYYMMDD\Orion\Salidas\YYYYMMDD_Gestion_orion.xlsx
    """
    from pathlib import Path

    base = Path(base)

    actual = base / "Orion" / "Salidas" / f"{fecha}_Gestion_orion.xlsx"

    return actual

# === CONSOLIDAR_V2_DATA_DIR_RUTAS_FIX_END ===


try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


RESUMEN_SQL = """
SELECT
    (SELECT COUNT_BIG(*) FROM Aster_Api.dbo.aster_dia_nc
     WHERE Fecha_Hora >= ? AND Fecha_Hora < ?) AS Total_Aster_dia_nc,

    (SELECT COUNT_BIG(*) FROM Aster_Api.dbo.comentarios
     WHERE fecha >= ? AND fecha < ?) AS Total_Aster_comentarios,

    (SELECT COUNT_BIG(*) FROM Aster_Api.dbo.usuarios) AS Total_Aster_Usuarios,

    (SELECT COUNT_BIG(*) FROM Orion.dbo.Causales
     WHERE FechayHora >= ? AND FechayHora < ?) AS Total_Orion_Causales,

    (SELECT COUNT_BIG(*) FROM Orion.dbo.Lote
     WHERE Fecha >= ? AND Fecha < ?) AS Total_Orion_Lote,

    (SELECT COUNT_BIG(*) FROM Orion.dbo.Discador
     WHERE FechayHora >= ? AND FechayHora < ?) AS Total_Orion_Discador;
"""

FASES = [
    ("A", "Preparar archivos"),
    ("B", "Unir archivos"),
    ("C", "Verificar calidad"),
    ("D", "Ajuste No contestan"),
    ("E", "Limpiar nota"),
    ("F", "Compromiso"),
    ("G", "Generar archivo final"),
    ("H", "Archivos generados"),
    ("I", "Cargar información"),
]


def normalizar_fecha_yyyymmdd(fecha_raw: str) -> str:
    limpia = re.sub(r"\D", "", str(fecha_raw or ""))

    if len(limpia) >= 8:
        return limpia[:8]

    raise ValueError(f"Fecha inválida: {fecha_raw}")


def fecha_iso(fecha_raw: str) -> str:
    fecha = normalizar_fecha_yyyymmdd(fecha_raw)
    return f"{fecha[:4]}-{fecha[4:6]}-{fecha[6:8]}"


def fecha_mes(fecha_raw: str) -> str:
    return normalizar_fecha_yyyymmdd(fecha_raw)[:6]


def resolver_rutas(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    fecha = normalizar_fecha_yyyymmdd(fecha_raw)
    base = Path(data_dir) / fecha

    ruta_orion = _gc_v2_resolver_archivo_orion_gestion(base, fecha)
    ruta_aster = _gc_v2_resolver_archivo_aster_gestion(base, fecha)

    carpeta_consolidado = base / "Consolidado" / "Gestion"

    subcarpetas = {
        "base": carpeta_consolidado,
        "00_resumen_sql": carpeta_consolidado / "00_resumen_sql",
        "01_union": carpeta_consolidado / "01_union",
        "02_verificaciones": carpeta_consolidado / "02_verificaciones",
        "03_ajustes": carpeta_consolidado / "03_ajustes",
        "04_compromiso": carpeta_consolidado / "04_compromiso",
        "05_final": carpeta_consolidado / "05_final",
        "06_carga": carpeta_consolidado / "06_carga",
        "Reportes": carpeta_consolidado / "Reportes",
        "Logs": carpeta_consolidado / "Logs",
    }

    return {
        "fecha": fecha,
        "fecha_iso": fecha_iso(fecha),
        "mes_gestion": fecha_mes(fecha),
        "base": base,
        "orion": ruta_orion,
        "aster": ruta_aster,
        "consolidado": carpeta_consolidado,
        "subcarpetas": subcarpetas,
    }


def crear_estructura_consolidado(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    rutas = resolver_rutas(data_dir, fecha_raw)

    creadas = []

    for nombre, ruta in rutas["subcarpetas"].items():
        ruta.mkdir(parents=True, exist_ok=True)
        creadas.append(
            {
                "carpeta": nombre,
                "ruta": str(ruta),
                "existe": ruta.exists(),
            }
        )

    return {
        "ok": True,
        "fecha": rutas["fecha"],
        "consolidado": str(rutas["consolidado"]),
        "carpetas": creadas,
    }


def contar_filas_excel(ruta: Path) -> int | None:
    if not ruta.exists() or ruta.suffix.lower() != ".xlsx":
        return None

    if load_workbook is None:
        return None

    try:
        wb = load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.active
        total = max(0, int(ws.max_row or 0) - 1)
        wb.close()
        return total
    except Exception:
        return None


def info_archivo(ruta: Path) -> dict[str, Any]:
    existe = ruta.exists()

    info = {
        "nombre": ruta.name,
        "ruta": str(ruta),
        "existe": existe,
        "tipo": ruta.suffix.lower().replace(".", "") if ruta.suffix else "",
        "tamano_kb": "",
        "fecha_creacion": "",
        "fecha_modificacion": "",
        "filas": "",
    }

    if not existe:
        return info

    stat = ruta.stat()

    info["tamano_kb"] = round(stat.st_size / 1024, 2)
    info["fecha_creacion"] = datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M:%S")
    info["fecha_modificacion"] = datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")

    filas = contar_filas_excel(ruta)

    if filas is not None:
        info["filas"] = filas

    return info


def preparar_proceso(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    rutas = resolver_rutas(data_dir, fecha_raw)
    estructura = crear_estructura_consolidado(data_dir, fecha_raw)

    archivos = [
        {
            "origen": "ORION",
            **info_archivo(rutas["orion"]),
        },
        {
            "origen": "ASTER",
            **info_archivo(rutas["aster"]),
        },
    ]

    ok_archivos = all(item["existe"] for item in archivos)

    return {
        "ok": ok_archivos,
        "fecha": rutas["fecha"],
        "fecha_iso": rutas["fecha_iso"],
        "mes_gestion": rutas["mes_gestion"],
        "conexion_requerida": "LOCAL/REMOTO",
        "carpeta_consolidado": str(rutas["consolidado"]),
        "archivos": archivos,
        "carpetas": estructura["carpetas"],
        "mensaje": (
            "Archivos base localizados correctamente."
            if ok_archivos
            else "Faltan archivos base ORION/ASTER. Revise rutas antes de continuar."
        ),
    }


def _connection_string(conexion: str) -> str:
    """
    Conexión SQL Server para resumen inicial.

    Variables opcionales:
    - GESTION_SQL_DRIVER
    - GESTION_SQL_LOCAL_SERVER
    - GESTION_SQL_REMOTE_SERVER
    - GESTION_SQL_USER
    - GESTION_SQL_PASSWORD

    Si no hay usuario/clave, usa Trusted_Connection=yes.
    """
    conn = str(conexion or "local").lower()
    driver = os.getenv("GESTION_SQL_DRIVER", "ODBC Driver 17 for SQL Server")

    local_server = os.getenv("GESTION_SQL_LOCAL_SERVER", r"localhost\SQL2025DEV")
    remote_server = os.getenv("GESTION_SQL_REMOTE_SERVER", "VC-EIDER")

    server = remote_server if conn == "remoto" else local_server

    user = os.getenv("GESTION_SQL_USER") or os.getenv("SQLSERVER_USER")
    password = os.getenv("GESTION_SQL_PASSWORD") or os.getenv("SQLSERVER_PASSWORD")

    if user and password:
        return (
            f"DRIVER={{{driver}}};"
            f"SERVER={server};"
            "DATABASE=Aster_Api;"
            f"UID={user};"
            f"PWD={password};"
            "TrustServerCertificate=yes;"
        )

    return (
        f"DRIVER={{{driver}}};"
        f"SERVER={server};"
        "DATABASE=Aster_Api;"
        "Trusted_Connection=yes;"
        "TrustServerCertificate=yes;"
    )


def consultar_resumen_sql(fecha_raw: str, conexion: str) -> dict[str, Any]:
    try:
        import pyodbc
    except Exception as exc:
        return {
            "ok": False,
            "conexion": conexion,
            "fecha": fecha_iso(fecha_raw),
            "error": f"No se pudo importar pyodbc: {exc}",
            "datos": {},
        }

    fecha = fecha_iso(fecha_raw)

    from datetime import datetime as _dt_resumen_sql, timedelta as _td_resumen_sql

    fecha_fin = (
        _dt_resumen_sql.strptime(fecha, "%Y-%m-%d")
        + _td_resumen_sql(days=1)
    ).strftime("%Y-%m-%d")

    params = [
        fecha, fecha_fin,
        fecha, fecha_fin,
        fecha, fecha_fin,
        fecha, fecha_fin,
        fecha, fecha_fin,
    ]

    try:
        with pyodbc.connect(_connection_string(conexion), timeout=15) as conn:
            cursor = conn.cursor()
            row = cursor.execute(RESUMEN_SQL, params).fetchone()

            if row is None:
                datos = {}
            else:
                columns = [col[0] for col in cursor.description]
                datos = dict(zip(columns, row))

        return {
            "ok": True,
            "conexion": str(conexion or "local").lower(),
            "fecha": fecha,
            "datos": datos,
            "error": "",
        }
    except Exception as exc:
        return {
            "ok": False,
            "conexion": str(conexion or "local").lower(),
            "fecha": fecha,
            "datos": {},
            "error": str(exc),
        }



def _normalizar_columnas_gestion(df: Any, origen: str) -> Any:
    df = df.copy()
    df.columns = [str(col).strip() for col in df.columns]

    if origen.upper() == "ORION" and "NroCliente_Contrato" in df.columns:
        df = df.rename(columns={"NroCliente_Contrato": "Cliente Nro."})

    if "Cliente Nro." not in df.columns:
        raise ValueError(
            f"El archivo {origen} no tiene la columna requerida Cliente Nro. "
            f"Columnas disponibles: {list(df.columns)}"
        )

    df["Cliente Nro."] = df["Cliente Nro."].astype(str).str.strip()

    return df


def _comparar_encabezados_gestion(cols_aster: list[str], cols_orion: list[str]) -> dict[str, Any]:
    set_aster = set(cols_aster)
    set_orion = set(cols_orion)

    faltan_en_orion = [col for col in cols_aster if col not in set_orion]
    sobran_en_orion = [col for col in cols_orion if col not in set_aster]
    mismo_orden = cols_aster == cols_orion
    mismo_set = not faltan_en_orion and not sobran_en_orion

    return {
        "ok": mismo_set,
        "mismo_orden": mismo_orden,
        "total_columnas_aster": len(cols_aster),
        "total_columnas_orion": len(cols_orion),
        "columnas_aster": cols_aster,
        "columnas_orion": cols_orion,
        "faltan_en_orion": faltan_en_orion,
        "sobran_en_orion": sobran_en_orion,
    }


def unir_archivos_gestion(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    """
    Fase B - Unión de archivos ASTER + ORION.

    Genera:
    - 01_union/YYYYMMDD_Gestion_union.xlsx
    - Reportes/YYYYMMDD_reporte_union.xlsx
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]

    ruta_aster = Path(rutas["aster"])
    ruta_orion = Path(rutas["orion"])

    errores = []

    if not ruta_aster.exists():
        errores.append(f"No existe archivo ASTER: {ruta_aster}")

    if not ruta_orion.exists():
        errores.append(f"No existe archivo ORION: {ruta_orion}")

    if errores:
        return {
            "ok": False,
            "fecha": fecha,
            "error": " | ".join(errores),
            "ruta_aster": str(ruta_aster),
            "ruta_orion": str(ruta_orion),
        }

    try:
        df_aster_raw = pd.read_excel(ruta_aster)
        df_orion_raw = pd.read_excel(ruta_orion)

        df_aster = _normalizar_columnas_gestion(df_aster_raw, "ASTER")
        df_orion = _normalizar_columnas_gestion(df_orion_raw, "ORION")

        encabezados = _comparar_encabezados_gestion(
            list(df_aster.columns),
            list(df_orion.columns),
        )

        if not encabezados["ok"]:
            return {
                "ok": False,
                "fecha": fecha,
                "error": "Los encabezados de ASTER y ORION no son compatibles.",
                "encabezados": encabezados,
                "ruta_aster": str(ruta_aster),
                "ruta_orion": str(ruta_orion),
                "total_aster": len(df_aster),
                "total_orion": len(df_orion),
            }

        if not encabezados["mismo_orden"]:
            df_orion = df_orion[list(df_aster.columns)]

        total_aster = len(df_aster)
        total_orion = len(df_orion)

        df_union = pd.concat([df_aster, df_orion], ignore_index=True)
        total_union = len(df_union)

        carpeta_union = Path(rutas["subcarpetas"]["01_union"])
        carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])
        carpeta_union.mkdir(parents=True, exist_ok=True)
        carpeta_reportes.mkdir(parents=True, exist_ok=True)

        ruta_union = carpeta_union / f"{fecha}_Gestion_union.xlsx"
        ruta_reporte = carpeta_reportes / f"{fecha}_reporte_union.xlsx"

        df_union.to_excel(ruta_union, index=False)

        resumen = pd.DataFrame(
            [
                {"concepto": "Total ASTER", "valor": total_aster},
                {"concepto": "Total ORION", "valor": total_orion},
                {"concepto": "Total después de unión", "valor": total_union},
                {"concepto": "Columnas ASTER", "valor": encabezados["total_columnas_aster"]},
                {"concepto": "Columnas ORION", "valor": encabezados["total_columnas_orion"]},
                {"concepto": "Encabezados compatibles", "valor": "SI"},
                {"concepto": "Mismo orden de columnas", "valor": "SI" if encabezados["mismo_orden"] else "NO - ORION fue reordenado"},
                {"concepto": "Archivo unión", "valor": str(ruta_union)},
            ]
        )

        columnas = pd.DataFrame(
            {
                "orden": list(range(1, len(df_aster.columns) + 1)),
                "columna_final": list(df_aster.columns),
            }
        )

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            columnas.to_excel(writer, sheet_name="Columnas", index=False)

        return {
            "ok": True,
            "fecha": fecha,
            "ruta_aster": str(ruta_aster),
            "ruta_orion": str(ruta_orion),
            "ruta_union": str(ruta_union),
            "ruta_reporte": str(ruta_reporte),
            "archivo_union": ruta_union.name,
            "archivo_reporte": ruta_reporte.name,
            "total_aster": total_aster,
            "total_orion": total_orion,
            "total_union": total_union,
            "encabezados": encabezados,
            "columnas": list(df_aster.columns),
        }
    except Exception as exc:
        return {
            "ok": False,
            "fecha": fecha,
            "error": str(exc),
            "ruta_aster": str(ruta_aster),
            "ruta_orion": str(ruta_orion),
        }



def _normalizar_nombre_columna_gestion(valor: str) -> str:
    import unicodedata

    texto = str(valor or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    texto = re.sub(r"\s+", " ", texto).strip()

    return texto


def _buscar_columna_gestion(df: Any, candidatos: list[str]) -> str:
    mapa = {
        _normalizar_nombre_columna_gestion(col): col
        for col in df.columns
    }

    for candidato in candidatos:
        key = _normalizar_nombre_columna_gestion(candidato)

        if key in mapa:
            return mapa[key]

    disponibles = ", ".join(str(col) for col in df.columns)

    raise ValueError(
        f"No se encontró ninguna columna candidata {candidatos}. "
        f"Columnas disponibles: {disponibles}"
    )


def _serie_texto_gestion(df: Any, columna: str) -> Any:
    return df[columna].fillna("").astype(str).str.strip()


def verificar_calidad_gestion(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    """
    Fase C - Verificaciones de calidad.

    Verifica:
    C1. Descripcion Codigo De Gestion: valores únicos y vacíos.
    C2. Clase de Gestion contiene TEL y debe tener Asesor/Grabador.
    C3. Duplicidad TEL por Cliente Nro.; conserva el primer registro por defecto.

    Genera:
    - 02_verificaciones/YYYYMMDD_Gestion_verificada.xlsx
    - Reportes/YYYYMMDD_reporte_verificacion_calidad.xlsx
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]

    ruta_union = Path(rutas["subcarpetas"]["01_union"]) / f"{fecha}_Gestion_union.xlsx"

    if not ruta_union.exists():
        return {
            "ok": False,
            "error": f"No existe el archivo de unión. Ejecute primero Fase B: {ruta_union}",
            "ruta_union": str(ruta_union),
        }

    try:
        df = pd.read_excel(ruta_union)
        df.columns = [str(col).strip() for col in df.columns]

        total_inicial = len(df)

        col_cliente = _buscar_columna_gestion(
            df,
            ["Cliente Nro.", "Cliente Nro", "NroCliente_Contrato", "Codigo_Cliente", "Código Cliente"],
        )
        col_desc = _buscar_columna_gestion(
            df,
            ["Descripcion Codigo De Gestion", "Descripción Código De Gestión", "Descripcion Codigo Gestion"],
        )
        col_clase = _buscar_columna_gestion(
            df,
            ["Clase de Gestion", "Clase de Gestión"],
        )
        col_asesor = _buscar_columna_gestion(
            df,
            ["Asesor"],
        )
        col_grabador = _buscar_columna_gestion(
            df,
            ["Grabador"],
        )

        df[col_cliente] = _serie_texto_gestion(df, col_cliente)

        serie_desc = _serie_texto_gestion(df, col_desc)
        serie_clase = _serie_texto_gestion(df, col_clase)
        serie_asesor = _serie_texto_gestion(df, col_asesor)
        serie_grabador = _serie_texto_gestion(df, col_grabador)

        # C1 - Valores únicos e inconsistencias vacías
        valores_unicos = (
            serie_desc.replace("", "(VACÍO)")
            .value_counts(dropna=False)
            .reset_index()
        )
        valores_unicos.columns = ["Descripcion Codigo De Gestion", "total"]
        valores_unicos["seleccionado_por_defecto"] = "SI"

        mask_desc_vacia = serie_desc.eq("")
        df_desc_vacia = df.loc[mask_desc_vacia].copy()

        # C2 - TEL sin Asesor/Grabador
        mask_tel = serie_clase.str.contains("TEL", case=False, na=False)
        mask_sin_asesor = serie_asesor.eq("")
        mask_sin_grabador = serie_grabador.eq("")
        mask_tel_incompleto = mask_tel & (mask_sin_asesor | mask_sin_grabador)

        df_tel_incompleto = df.loc[mask_tel_incompleto].copy()

        # C3 - Duplicidad TEL por Cliente Nro.
        serie_cliente = _serie_texto_gestion(df, col_cliente)
        mask_dup_tel_all = mask_tel & serie_cliente.duplicated(keep=False)
        mask_dup_tel_eliminar = mask_tel & serie_cliente.duplicated(keep="first")

        df_tel_duplicados = df.loc[mask_dup_tel_all].copy()
        df_tel_eliminados = df.loc[mask_dup_tel_eliminar].copy()

        df_verificada = df.drop(index=df_tel_eliminados.index).copy().reset_index(drop=True)
        total_final = len(df_verificada)

        carpeta_verificacion = Path(rutas["subcarpetas"]["02_verificaciones"])
        carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])
        carpeta_verificacion.mkdir(parents=True, exist_ok=True)
        carpeta_reportes.mkdir(parents=True, exist_ok=True)

        ruta_verificada = carpeta_verificacion / f"{fecha}_Gestion_verificada.xlsx"
        ruta_reporte = carpeta_reportes / f"{fecha}_reporte_verificacion_calidad.xlsx"

        ruta_desc_vacia = carpeta_reportes / f"{fecha}_inconsistencias_descripcion_vacia.xlsx"
        ruta_tel_incompleto = carpeta_reportes / f"{fecha}_tel_sin_asesor_grabador.xlsx"
        ruta_tel_eliminados = carpeta_reportes / f"{fecha}_tel_duplicados_eliminados.xlsx"

        df_verificada.to_excel(ruta_verificada, index=False)

        if len(df_desc_vacia):
            df_desc_vacia.to_excel(ruta_desc_vacia, index=False)

        if len(df_tel_incompleto):
            df_tel_incompleto.to_excel(ruta_tel_incompleto, index=False)

        if len(df_tel_eliminados):
            df_tel_eliminados.to_excel(ruta_tel_eliminados, index=False)

        resumen = pd.DataFrame(
            [
                {"control": "Total inicial", "valor": total_inicial},
                {"control": "Valores únicos Descripcion Codigo De Gestion", "valor": len(valores_unicos)},
                {"control": "Descripcion vacía", "valor": len(df_desc_vacia)},
                {"control": "Registros TEL", "valor": int(mask_tel.sum())},
                {"control": "TEL sin Asesor o Grabador", "valor": len(df_tel_incompleto)},
                {"control": "Duplicados TEL detectados", "valor": len(df_tel_duplicados)},
                {"control": "Duplicados TEL eliminados por defecto", "valor": len(df_tel_eliminados)},
                {"control": "Total final verificado", "valor": total_final},
            ]
        )

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            valores_unicos.to_excel(writer, sheet_name="ValoresDescripcion", index=False)
            df_desc_vacia.to_excel(writer, sheet_name="DescVacia", index=False)
            df_tel_incompleto.to_excel(writer, sheet_name="TelSinAsesorGrabador", index=False)
            df_tel_duplicados.to_excel(writer, sheet_name="TelDuplicados", index=False)
            df_tel_eliminados.to_excel(writer, sheet_name="TelEliminados", index=False)

        requiere_revision = any(
            [
                len(df_desc_vacia) > 0,
                len(df_tel_incompleto) > 0,
                len(df_tel_eliminados) > 0,
            ]
        )

        return {
            "ok": True,
            "requiere_revision": requiere_revision,
            "fecha": fecha,
            "ruta_union": str(ruta_union),
            "ruta_verificada": str(ruta_verificada),
            "ruta_reporte": str(ruta_reporte),
            "ruta_desc_vacia": str(ruta_desc_vacia) if len(df_desc_vacia) else "",
            "ruta_tel_incompleto": str(ruta_tel_incompleto) if len(df_tel_incompleto) else "",
            "ruta_tel_eliminados": str(ruta_tel_eliminados) if len(df_tel_eliminados) else "",
            "columnas": {
                "cliente": col_cliente,
                "descripcion": col_desc,
                "clase": col_clase,
                "asesor": col_asesor,
                "grabador": col_grabador,
            },
            "totales": {
                "total_inicial": total_inicial,
                "valores_unicos_descripcion": len(valores_unicos),
                "descripcion_vacia": len(df_desc_vacia),
                "registros_tel": int(mask_tel.sum()),
                "tel_sin_asesor_grabador": len(df_tel_incompleto),
                "tel_duplicados_detectados": len(df_tel_duplicados),
                "tel_duplicados_eliminados": len(df_tel_eliminados),
                "total_final": total_final,
            },
            "valores_unicos": valores_unicos.head(300).to_dict(orient="records"),
            "tel_incompleto_preview": df_tel_incompleto.head(80).to_dict(orient="records"),
            "tel_duplicados_preview": df_tel_duplicados.head(80).to_dict(orient="records"),
            "tel_eliminados_preview": df_tel_eliminados.head(80).to_dict(orient="records"),
        }
    except Exception as exc:
        return {
            "ok": False,
            "fecha": fecha,
            "error": str(exc),
            "ruta_union": str(ruta_union),
        }



def _porcentaje_no_contestan_valido(valor: Any) -> float:
    try:
        pct = float(str(valor or "12").replace(",", "."))
    except Exception:
        pct = 12.0

    if pct < 10:
        pct = 10.0

    if pct > 15:
        pct = 15.0

    return pct


def aplicar_ajuste_no_contestan_gestion(
    data_dir: str | Path,
    fecha_raw: str,
    porcentaje_raw: Any = 12,
) -> dict[str, Any]:
    """
    Fase D - Ajuste No contestan.

    Para Tipo Cartera Home/Mobile, reemplaza entre 10% y 15% de:
    - Buzon de voz
    - Telefono Fuera de Servicio
    - Telefono Ocupado

    por:
    - No contestan

    Genera:
    - 03_ajustes/YYYYMMDD_Gestion_ajuste_no_contestan.xlsx
    - Reportes/YYYYMMDD_reporte_ajuste_no_contestan.xlsx
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]
    porcentaje = _porcentaje_no_contestan_valido(porcentaje_raw)

    ruta_entrada = (
        Path(rutas["subcarpetas"]["02_verificaciones"])
        / f"{fecha}_Gestion_verificada.xlsx"
    )

    if not ruta_entrada.exists():
        return {
            "ok": False,
            "error": f"No existe archivo verificado. Ejecute primero Fase C: {ruta_entrada}",
            "ruta_entrada": str(ruta_entrada),
        }

    try:
        df = pd.read_excel(ruta_entrada)
        df.columns = [str(col).strip() for col in df.columns]

        total_inicial = len(df)

        col_desc = _buscar_columna_gestion(
            df,
            [
                "Descripcion Codigo De Gestion",
                "Descripción Código De Gestión",
                "Descripcion Codigo Gestion",
            ],
        )
        col_cartera = _buscar_columna_gestion(
            df,
            [
                "Tipo Cartera",
                "Tipo_Cartera",
                "Cartera",
            ],
        )

        valores_objetivo = [
            "Buzon de voz",
            "Telefono Fuera de Servicio",
            "Telefono Ocupado",
        ]

        carteras = ["Home", "Mobile"]

        df_ajustado = df.copy()
        serie_desc = _serie_texto_gestion(df_ajustado, col_desc)
        serie_cartera = _serie_texto_gestion(df_ajustado, col_cartera)

        reporte = []
        indices_reemplazados = []

        for cartera in carteras:
            for descripcion in valores_objetivo:
                mask = (
                    serie_cartera.str.lower().eq(cartera.lower())
                    & serie_desc.str.lower().eq(descripcion.lower())
                )

                indices = list(df_ajustado.index[mask])
                total_encontrado = len(indices)

                if total_encontrado:
                    total_reemplazar = int(round(total_encontrado * porcentaje / 100))

                    if total_reemplazar <= 0:
                        total_reemplazar = 1

                    if total_reemplazar > total_encontrado:
                        total_reemplazar = total_encontrado

                    # Selección determinística para que el proceso sea repetible.
                    seed = int(fecha[-4:]) + len(cartera) + len(descripcion)
                    seleccion = (
                        df_ajustado.loc[indices]
                        .sample(n=total_reemplazar, random_state=seed)
                        .index
                        .tolist()
                    )
                else:
                    total_reemplazar = 0
                    seleccion = []

                if seleccion:
                    df_ajustado.loc[seleccion, col_desc] = "No contestan"
                    indices_reemplazados.extend(seleccion)

                reporte.append(
                    {
                        "Tipo Cartera": cartera,
                        "Descripcion original": descripcion,
                        "Total encontrados": total_encontrado,
                        "Porcentaje aplicado": porcentaje,
                        "Total reemplazados": len(seleccion),
                        "Nuevo valor": "No contestan",
                    }
                )

        total_final = len(df_ajustado)
        control_filas_ok = total_inicial == total_final

        carpeta_ajustes = Path(rutas["subcarpetas"]["03_ajustes"])
        carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])
        carpeta_ajustes.mkdir(parents=True, exist_ok=True)
        carpeta_reportes.mkdir(parents=True, exist_ok=True)

        ruta_salida = carpeta_ajustes / f"{fecha}_Gestion_ajuste_no_contestan.xlsx"
        ruta_reporte = carpeta_reportes / f"{fecha}_reporte_ajuste_no_contestan.xlsx"

        df_ajustado.to_excel(ruta_salida, index=False)

        df_reporte = pd.DataFrame(reporte)

        if indices_reemplazados:
            df_reemplazados = df.loc[sorted(set(indices_reemplazados))].copy()
            df_reemplazados["Nuevo valor"] = "No contestan"
        else:
            df_reemplazados = pd.DataFrame()

        resumen = pd.DataFrame(
            [
                {"control": "Total filas antes", "valor": total_inicial},
                {"control": "Total filas después", "valor": total_final},
                {"control": "Control filas iguales", "valor": "SI" if control_filas_ok else "NO"},
                {"control": "Porcentaje aplicado", "valor": porcentaje},
                {"control": "Total reemplazos", "valor": len(indices_reemplazados)},
                {"control": "Archivo salida", "valor": str(ruta_salida)},
            ]
        )

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            df_reporte.to_excel(writer, sheet_name="Detalle", index=False)
            df_reemplazados.to_excel(writer, sheet_name="Reemplazados", index=False)

        return {
            "ok": control_filas_ok,
            "fecha": fecha,
            "porcentaje": porcentaje,
            "ruta_entrada": str(ruta_entrada),
            "ruta_salida": str(ruta_salida),
            "ruta_reporte": str(ruta_reporte),
            "archivo_salida": ruta_salida.name,
            "archivo_reporte": ruta_reporte.name,
            "total_inicial": total_inicial,
            "total_final": total_final,
            "control_filas_ok": control_filas_ok,
            "total_reemplazos": len(indices_reemplazados),
            "detalle": reporte,
            "reemplazados_preview": df_reemplazados.head(80).to_dict(orient="records") if len(df_reemplazados) else [],
            "columnas": {
                "descripcion": col_desc,
                "tipo_cartera": col_cartera,
            },
        }
    except Exception as exc:
        return {
            "ok": False,
            "fecha": fecha,
            "error": str(exc),
            "ruta_entrada": str(ruta_entrada),
        }



def limpiar_nota_gestion(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    """
    Fase E - Limpieza columna Nota de la Gestion.

    Reemplaza:
    - "Acepta táctico de reconexión."
    por:
    - "Acepta táctico de reconexión"

    Genera:
    - 03_ajustes/YYYYMMDD_Gestion_nota_limpia.xlsx
    - Reportes/YYYYMMDD_reporte_limpieza_nota.xlsx
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]

    ruta_entrada = (
        Path(rutas["subcarpetas"]["03_ajustes"])
        / f"{fecha}_Gestion_ajuste_no_contestan.xlsx"
    )

    if not ruta_entrada.exists():
        return {
            "ok": False,
            "error": f"No existe archivo de Fase D. Ejecute primero Fase D: {ruta_entrada}",
            "ruta_entrada": str(ruta_entrada),
        }

    try:
        df = pd.read_excel(ruta_entrada)
        df.columns = [str(col).strip() for col in df.columns]

        total_inicial = len(df)

        col_nota = _buscar_columna_gestion(
            df,
            [
                "Nota de la Gestion",
                "Nota de la Gestión",
                "Nota Gestion",
                "Nota Gestión",
                "Nota",
            ],
        )

        valor_original = "Acepta táctico de reconexión."
        valor_nuevo = "Acepta táctico de reconexión"

        df_limpio = df.copy()
        serie_original = df_limpio[col_nota].fillna("").astype(str)

        mask_reemplazo = serie_original.str.contains(valor_original, regex=False, na=False)
        total_reemplazos = int(mask_reemplazo.sum())

        df_preview = df_limpio.loc[mask_reemplazo].copy()

        if total_reemplazos:
            df_preview["Nota antes"] = df_preview[col_nota].astype(str)

            df_limpio.loc[mask_reemplazo, col_nota] = (
                df_limpio.loc[mask_reemplazo, col_nota]
                .astype(str)
                .str.replace(valor_original, valor_nuevo, regex=False)
            )

            df_preview["Nota después"] = df_limpio.loc[mask_reemplazo, col_nota].astype(str)

        total_final = len(df_limpio)
        control_filas_ok = total_inicial == total_final

        carpeta_ajustes = Path(rutas["subcarpetas"]["03_ajustes"])
        carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])
        carpeta_ajustes.mkdir(parents=True, exist_ok=True)
        carpeta_reportes.mkdir(parents=True, exist_ok=True)

        ruta_salida = carpeta_ajustes / f"{fecha}_Gestion_nota_limpia.xlsx"
        ruta_reporte = carpeta_reportes / f"{fecha}_reporte_limpieza_nota.xlsx"

        df_limpio.to_excel(ruta_salida, index=False)

        resumen = pd.DataFrame(
            [
                {"control": "Total filas antes", "valor": total_inicial},
                {"control": "Total filas después", "valor": total_final},
                {"control": "Control filas iguales", "valor": "SI" if control_filas_ok else "NO"},
                {"control": "Columna usada", "valor": col_nota},
                {"control": "Texto buscado", "valor": valor_original},
                {"control": "Texto reemplazo", "valor": valor_nuevo},
                {"control": "Total reemplazos", "valor": total_reemplazos},
                {"control": "Archivo salida", "valor": str(ruta_salida)},
            ]
        )

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            df_preview.to_excel(writer, sheet_name="RegistrosLimpiados", index=False)

        return {
            "ok": control_filas_ok,
            "fecha": fecha,
            "ruta_entrada": str(ruta_entrada),
            "ruta_salida": str(ruta_salida),
            "ruta_reporte": str(ruta_reporte),
            "archivo_salida": ruta_salida.name,
            "archivo_reporte": ruta_reporte.name,
            "columna_nota": col_nota,
            "valor_original": valor_original,
            "valor_nuevo": valor_nuevo,
            "total_inicial": total_inicial,
            "total_final": total_final,
            "control_filas_ok": control_filas_ok,
            "total_reemplazos": total_reemplazos,
            "preview": df_preview.head(80).to_dict(orient="records") if total_reemplazos else [],
        }
    except Exception as exc:
        return {
            "ok": False,
            "fecha": fecha,
            "error": str(exc),
            "ruta_entrada": str(ruta_entrada),
        }



def procesar_compromiso_gestion(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    """
    Fase F - Compromiso.

    Reglas:
    - Filtra Descripcion Codigo De Gestion que inicia con "Acuerdo De Pago".
    - Si Fecha_Compromiso está vacía, cambia Descripcion Codigo De Gestion a:
      "No Hubo Acuerdo de pago".

    Genera:
    - 04_compromiso/YYYYMMDD_Gestion_compromiso.xlsx
    - 04_compromiso/YYYYMMDD_Gestion.csv
    - Reportes/YYYYMMDD_reporte_compromiso.xlsx
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]

    ruta_entrada = (
        Path(rutas["subcarpetas"]["03_ajustes"])
        / f"{fecha}_Gestion_nota_limpia.xlsx"
    )

    if not ruta_entrada.exists():
        return {
            "ok": False,
            "error": f"No existe archivo de Fase E. Ejecute primero Fase E: {ruta_entrada}",
            "ruta_entrada": str(ruta_entrada),
        }

    try:
        df = pd.read_excel(ruta_entrada)
        df.columns = [str(col).strip() for col in df.columns]

        total_inicial = len(df)

        col_desc = _buscar_columna_gestion(
            df,
            [
                "Descripcion Codigo De Gestion",
                "Descripción Código De Gestión",
                "Descripcion Codigo Gestion",
            ],
        )

        col_fecha_compromiso = _buscar_columna_gestion(
            df,
            [
                "Fecha_Compromiso",
                "Fecha Compromiso",
                "Fecha de Compromiso",
                "FechaCompromiso",
            ],
        )

        df_compromiso = df.copy()

        serie_desc = df_compromiso[col_desc].fillna("").astype(str).str.strip()
        serie_fecha = df_compromiso[col_fecha_compromiso]

        mask_acuerdo = serie_desc.str.lower().str.startswith("acuerdo de pago", na=False)

        serie_fecha_texto = serie_fecha.fillna("").astype(str).str.strip()
        mask_fecha_vacia = (
            serie_fecha.isna()
            | serie_fecha_texto.eq("")
            | serie_fecha_texto.str.lower().isin(["nan", "nat", "none", "null"])
        )

        mask_reemplazo = mask_acuerdo & mask_fecha_vacia
        mask_con_fecha = mask_acuerdo & ~mask_fecha_vacia

        total_acuerdo = int(mask_acuerdo.sum())
        total_con_fecha = int(mask_con_fecha.sum())
        total_sin_fecha = int(mask_reemplazo.sum())

        df_reemplazados = df_compromiso.loc[mask_reemplazo].copy()

        if total_sin_fecha:
            df_reemplazados["Descripcion antes"] = df_reemplazados[col_desc].astype(str)
            df_compromiso.loc[mask_reemplazo, col_desc] = "No Hubo Acuerdo de pago"
            df_reemplazados["Descripcion después"] = "No Hubo Acuerdo de pago"

        total_final = len(df_compromiso)
        control_filas_ok = total_inicial == total_final

        carpeta_compromiso = Path(rutas["subcarpetas"]["04_compromiso"])
        carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])
        carpeta_compromiso.mkdir(parents=True, exist_ok=True)
        carpeta_reportes.mkdir(parents=True, exist_ok=True)

        ruta_excel = carpeta_compromiso / f"{fecha}_Gestion_compromiso.xlsx"
        ruta_csv = carpeta_compromiso / f"{fecha}_Gestion.csv"
        ruta_reporte = carpeta_reportes / f"{fecha}_reporte_compromiso.xlsx"

        df_compromiso.to_excel(ruta_excel, index=False)
        df_compromiso.to_csv(ruta_csv, index=False, encoding="utf-8-sig")

        resumen = pd.DataFrame(
            [
                {"control": "Total filas antes", "valor": total_inicial},
                {"control": "Total filas después", "valor": total_final},
                {"control": "Control filas iguales", "valor": "SI" if control_filas_ok else "NO"},
                {"control": "Columna descripción", "valor": col_desc},
                {"control": "Columna Fecha_Compromiso", "valor": col_fecha_compromiso},
                {"control": "Acuerdo De Pago total", "valor": total_acuerdo},
                {"control": "Acuerdo De Pago con Fecha_Compromiso", "valor": total_con_fecha},
                {"control": "Acuerdo De Pago sin Fecha_Compromiso", "valor": total_sin_fecha},
                {"control": "Reemplazados a No Hubo Acuerdo de pago", "valor": total_sin_fecha},
                {"control": "Archivo Excel compromiso", "valor": str(ruta_excel)},
                {"control": "Archivo CSV gestión", "valor": str(ruta_csv)},
            ]
        )

        df_acuerdos = df_compromiso.loc[mask_acuerdo].copy()

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            df_acuerdos.to_excel(writer, sheet_name="AcuerdosPago", index=False)
            df_reemplazados.to_excel(writer, sheet_name="Reemplazados", index=False)

        return {
            "ok": control_filas_ok,
            "fecha": fecha,
            "ruta_entrada": str(ruta_entrada),
            "ruta_excel": str(ruta_excel),
            "ruta_csv": str(ruta_csv),
            "ruta_reporte": str(ruta_reporte),
            "archivo_excel": ruta_excel.name,
            "archivo_csv": ruta_csv.name,
            "archivo_reporte": ruta_reporte.name,
            "columna_descripcion": col_desc,
            "columna_fecha_compromiso": col_fecha_compromiso,
            "total_inicial": total_inicial,
            "total_final": total_final,
            "control_filas_ok": control_filas_ok,
            "total_acuerdo": total_acuerdo,
            "total_con_fecha": total_con_fecha,
            "total_sin_fecha": total_sin_fecha,
            "total_reemplazados": total_sin_fecha,
            "preview_reemplazados": df_reemplazados.head(80).to_dict(orient="records") if total_sin_fecha else [],
        }
    except Exception as exc:
        return {
            "ok": False,
            "fecha": fecha,
            "error": str(exc),
            "ruta_entrada": str(ruta_entrada),
        }



def _limpiar_cliente_nro_gestion(valor: Any) -> str:
    """
    Limpia Cliente Nro.:
    - convierte a texto
    - quita espacios
    - elimina letras m/M según regla del proceso
    - elimina .0 cuando Excel lo leyó como número
    """
    texto = str(valor if valor is not None else "").strip()

    if texto.lower() in {"nan", "nat", "none", "null"}:
        texto = ""

    if texto.endswith(".0"):
        texto = texto[:-2]

    texto = texto.replace("m", "").replace("M", "").strip()

    return texto


def generar_archivo_final_gestion(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    """
    Fase G - Código cliente y archivo final.

    Reglas:
    - Limpia Cliente Nro. / NroCliente_Contrato.
    - Agrega Mes_Gestion.
    - Agrega origen_datos = ondedrive.
    - Agrega crm:
        orion si Cliente Nro. cruza con NroCliente_Contrato del archivo ORION original.
        aster si no cruza.
    - Exporta:
        05_final/YYYYMMDD_Gestion_excel.xlsx
        Reportes/YYYYMMDD_reporte_archivo_final.xlsx
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]
    mes_gestion = rutas["mes_gestion"]

    ruta_entrada = (
        Path(rutas["subcarpetas"]["04_compromiso"])
        / f"{fecha}_Gestion_compromiso.xlsx"
    )
    ruta_orion_original = Path(rutas["orion"])

    if not ruta_entrada.exists():
        return {
            "ok": False,
            "error": f"No existe archivo de Fase F. Ejecute primero Fase F: {ruta_entrada}",
            "ruta_entrada": str(ruta_entrada),
        }

    if not ruta_orion_original.exists():
        return {
            "ok": False,
            "error": f"No existe archivo ORION original para cruce CRM: {ruta_orion_original}",
            "ruta_entrada": str(ruta_entrada),
            "ruta_orion": str(ruta_orion_original),
        }

    try:
        df = pd.read_excel(ruta_entrada)
        df.columns = [str(col).strip() for col in df.columns]

        total_inicial = len(df)

        col_cliente = _buscar_columna_gestion(
            df,
            [
                "Cliente Nro.",
                "Cliente Nro",
                "NroCliente_Contrato",
                "Codigo_Cliente",
                "Código Cliente",
            ],
        )

        df_final = df.copy()

        clientes_antes_vacios = int(
            df_final[col_cliente].fillna("").astype(str).str.strip().eq("").sum()
        )

        df_final[col_cliente] = df_final[col_cliente].apply(_limpiar_cliente_nro_gestion)

        clientes_despues_vacios = int(df_final[col_cliente].eq("").sum())

        # Leer archivo ORION original y construir set de clientes ORION.
        df_orion = pd.read_excel(ruta_orion_original)
        df_orion.columns = [str(col).strip() for col in df_orion.columns]

        col_orion_cliente = _buscar_columna_gestion(
            df_orion,
            [
                "NroCliente_Contrato",
                "Cliente Nro.",
                "Cliente Nro",
                "Codigo_Cliente",
                "Código Cliente",
            ],
        )

        clientes_orion = set(
            df_orion[col_orion_cliente]
            .apply(_limpiar_cliente_nro_gestion)
            .dropna()
            .astype(str)
            .str.strip()
        )
        clientes_orion.discard("")

        df_final["Mes_Gestion"] = mes_gestion
        df_final["origen_datos"] = "ondedrive"
        df_final["crm"] = df_final[col_cliente].apply(
            lambda value: "orion" if str(value).strip() in clientes_orion else "aster"
        )

        total_orion = int((df_final["crm"] == "orion").sum())
        total_aster = int((df_final["crm"] == "aster").sum())

        total_final = len(df_final)
        control_filas_ok = total_inicial == total_final

        carpeta_final = Path(rutas["subcarpetas"]["05_final"])
        carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])
        carpeta_final.mkdir(parents=True, exist_ok=True)
        carpeta_reportes.mkdir(parents=True, exist_ok=True)

        ruta_excel = carpeta_final / f"{fecha}_Gestion_excel.xlsx"
        ruta_reporte = carpeta_reportes / f"{fecha}_reporte_archivo_final.xlsx"

        df_final.to_excel(ruta_excel, index=False)

        resumen = pd.DataFrame(
            [
                {"control": "Total filas antes", "valor": total_inicial},
                {"control": "Total filas después", "valor": total_final},
                {"control": "Control filas iguales", "valor": "SI" if control_filas_ok else "NO"},
                {"control": "Columna Cliente Nro.", "valor": col_cliente},
                {"control": "Columna ORION cruce", "valor": col_orion_cliente},
                {"control": "Mes_Gestion", "valor": mes_gestion},
                {"control": "origen_datos", "valor": "ondedrive"},
                {"control": "Total crm orion", "valor": total_orion},
                {"control": "Total crm aster", "valor": total_aster},
                {"control": "Clientes vacíos antes", "valor": clientes_antes_vacios},
                {"control": "Clientes vacíos después", "valor": clientes_despues_vacios},
                {"control": "Clientes únicos ORION para cruce", "valor": len(clientes_orion)},
                {"control": "Archivo final", "valor": str(ruta_excel)},
            ]
        )

        crm_resumen = pd.DataFrame(
            [
                {"crm": "orion", "total": total_orion},
                {"crm": "aster", "total": total_aster},
            ]
        )

        preview = df_final.head(120).copy()

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            crm_resumen.to_excel(writer, sheet_name="CRM", index=False)
            preview.to_excel(writer, sheet_name="PreviewFinal", index=False)

        return {
            "ok": control_filas_ok,
            "fecha": fecha,
            "mes_gestion": mes_gestion,
            "ruta_entrada": str(ruta_entrada),
            "ruta_orion": str(ruta_orion_original),
            "ruta_excel": str(ruta_excel),
            "ruta_reporte": str(ruta_reporte),
            "archivo_excel": ruta_excel.name,
            "archivo_reporte": ruta_reporte.name,
            "columna_cliente": col_cliente,
            "columna_orion_cliente": col_orion_cliente,
            "total_inicial": total_inicial,
            "total_final": total_final,
            "control_filas_ok": control_filas_ok,
            "total_orion": total_orion,
            "total_aster": total_aster,
            "clientes_antes_vacios": clientes_antes_vacios,
            "clientes_despues_vacios": clientes_despues_vacios,
            "clientes_orion_unicos": len(clientes_orion),
            "origen_datos": "ondedrive",
            "preview": preview.head(80).to_dict(orient="records"),
        }
    except Exception as exc:
        return {
            "ok": False,
            "fecha": fecha,
            "error": str(exc),
            "ruta_entrada": str(ruta_entrada),
            "ruta_orion": str(ruta_orion_original),
        }



def _contar_filas_archivo_generado(ruta: Path) -> int | str:
    ext = ruta.suffix.lower()

    try:
        if ext == ".xlsx":
            return contar_filas_excel(ruta) or ""

        if ext == ".csv":
            with ruta.open("r", encoding="utf-8-sig", errors="ignore") as fh:
                total = sum(1 for _ in fh)

            return max(0, total - 1)
    except Exception:
        return ""

    return ""


def _info_archivo_generado_gestion(ruta: Path, base_consolidado: Path) -> dict[str, Any]:
    stat = ruta.stat()
    ext = ruta.suffix.lower().replace(".", "")

    try:
        carpeta_relativa = str(ruta.parent.relative_to(base_consolidado))
    except Exception:
        carpeta_relativa = ruta.parent.name

    return {
        "nombre": ruta.name,
        "ruta": str(ruta),
        "tipo": ext.upper(),
        "carpeta": carpeta_relativa,
        "tamano_kb": round(stat.st_size / 1024, 2),
        "fecha_creacion": datetime.fromtimestamp(stat.st_ctime).strftime("%Y-%m-%d %H:%M:%S"),
        "fecha_modificacion": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "filas": _contar_filas_archivo_generado(ruta),
    }


def listar_archivos_generados_gestion(data_dir: str | Path, fecha_raw: str) -> dict[str, Any]:
    """
    Fase H - Archivos generados.

    Lista todos los .xlsx y .csv creados dentro de:
    data/YYYYMMDD/Consolidado/Gestion

    Genera además:
    - Reportes/YYYYMMDD_reporte_archivos_generados.xlsx
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]
    carpeta_consolidado = Path(rutas["consolidado"])
    carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])

    if not carpeta_consolidado.exists():
        return {
            "ok": False,
            "error": f"No existe la carpeta Consolidado/Gestion. Ejecute primero Fase A: {carpeta_consolidado}",
            "carpeta_consolidado": str(carpeta_consolidado),
        }

    try:
        archivos = []

        for patron in ("*.xlsx", "*.csv"):
            for ruta in carpeta_consolidado.rglob(patron):
                if ruta.is_file():
                    archivos.append(_info_archivo_generado_gestion(ruta, carpeta_consolidado))

        archivos = sorted(
            archivos,
            key=lambda item: (
                str(item.get("tipo", "")),
                str(item.get("carpeta", "")),
                str(item.get("nombre", "")),
            ),
        )

        total_archivos = len(archivos)
        total_excel = sum(1 for item in archivos if item.get("tipo") == "XLSX")
        total_csv = sum(1 for item in archivos if item.get("tipo") == "CSV")
        tamano_total_kb = round(sum(float(item.get("tamano_kb") or 0) for item in archivos), 2)

        archivo_final = (
            Path(rutas["subcarpetas"]["05_final"])
            / f"{fecha}_Gestion_excel.xlsx"
        )
        archivo_csv = (
            Path(rutas["subcarpetas"]["04_compromiso"])
            / f"{fecha}_Gestion.csv"
        )

        carpeta_reportes.mkdir(parents=True, exist_ok=True)
        ruta_reporte = carpeta_reportes / f"{fecha}_reporte_archivos_generados.xlsx"

        df_archivos = pd.DataFrame(archivos)

        resumen = pd.DataFrame(
            [
                {"control": "Total archivos", "valor": total_archivos},
                {"control": "Total Excel", "valor": total_excel},
                {"control": "Total CSV", "valor": total_csv},
                {"control": "Tamaño total KB", "valor": tamano_total_kb},
                {"control": "Archivo final existe", "valor": "SI" if archivo_final.exists() else "NO"},
                {"control": "Archivo CSV existe", "valor": "SI" if archivo_csv.exists() else "NO"},
                {"control": "Carpeta Consolidado/Gestion", "valor": str(carpeta_consolidado)},
            ]
        )

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            df_archivos.to_excel(writer, sheet_name="Archivos", index=False)

        # Incluir el reporte recién generado en el resultado visual.
        if ruta_reporte.exists():
            info_reporte = _info_archivo_generado_gestion(ruta_reporte, carpeta_consolidado)

            if not any(item.get("ruta") == str(ruta_reporte) for item in archivos):
                archivos.append(info_reporte)

        archivos = sorted(
            archivos,
            key=lambda item: (
                str(item.get("tipo", "")),
                str(item.get("carpeta", "")),
                str(item.get("nombre", "")),
            ),
        )

        return {
            "ok": True,
            "fecha": fecha,
            "carpeta_consolidado": str(carpeta_consolidado),
            "ruta_reporte": str(ruta_reporte),
            "archivo_reporte": ruta_reporte.name,
            "archivo_final": str(archivo_final),
            "archivo_csv": str(archivo_csv),
            "archivo_final_existe": archivo_final.exists(),
            "archivo_csv_existe": archivo_csv.exists(),
            "total_archivos": len(archivos),
            "total_excel": sum(1 for item in archivos if item.get("tipo") == "XLSX"),
            "total_csv": sum(1 for item in archivos if item.get("tipo") == "CSV"),
            "tamano_total_kb": round(sum(float(item.get("tamano_kb") or 0) for item in archivos), 2),
            "archivos": archivos,
        }
    except Exception as exc:
        return {
            "ok": False,
            "fecha": fecha,
            "error": str(exc),
            "carpeta_consolidado": str(carpeta_consolidado),
        }



def _gc_existe_fecha_gestion_sql(cursor: Any, schema: str, table: str, fecha_yyyymmdd: str) -> dict[str, Any]:
    """
    Verifica si ya existen registros para la fecha del proceso.

    Usa rango de fecha:
        Fecha_De_Gestion >= yyyy-mm-dd
        Fecha_De_Gestion <  yyyy-mm-dd + 1

    No usa CAST(Fecha_De_Gestion AS DATE) para permitir uso de índice si existe.
    """
    from datetime import datetime, timedelta

    fecha_limpia = normalizar_fecha_yyyymmdd(fecha_yyyymmdd)
    fecha_inicio_dt = datetime.strptime(fecha_limpia, "%Y%m%d")
    fecha_fin_dt = fecha_inicio_dt + timedelta(days=1)

    fecha_inicio = fecha_inicio_dt.strftime("%Y-%m-%d")
    fecha_fin = fecha_fin_dt.strftime("%Y-%m-%d")

    sql_existe = f"""
    SELECT TOP (1) 1
    FROM [{schema}].[{table}]
    WHERE [Fecha_De_Gestion] >= ?
      AND [Fecha_De_Gestion] < ?
    """

    row = cursor.execute(sql_existe, fecha_inicio, fecha_fin).fetchone()
    existe = row is not None

    return {
        "existe": existe,
        "fecha_columna": "Fecha_De_Gestion",
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "mensaje": (
            "Ya existen registros para la fecha del proceso. Inserción bloqueada para evitar duplicados."
            if existe
            else "No existen registros previos para la fecha del proceso."
        ),
    }

def _connection_string_vencorp(conexion: str) -> str:
    """
    Conexión SQL Server para carga final de Consolidar Gestión.

    Variables opcionales:
    - GESTION_SQL_DRIVER
    - GESTION_SQL_LOCAL_SERVER
    - GESTION_SQL_REMOTE_SERVER
    - GESTION_SQL_USER
    - GESTION_SQL_PASSWORD

    Si no hay usuario/clave, usa Trusted_Connection=yes.
    """
    conn = str(conexion or "local").lower()
    driver = os.getenv("GESTION_SQL_DRIVER", "ODBC Driver 17 for SQL Server")

    local_server = os.getenv("GESTION_SQL_LOCAL_SERVER", r"localhost\SQL2025DEV")
    remote_server = os.getenv("GESTION_SQL_REMOTE_SERVER", "VC-EIDER")

    server = remote_server if conn == "remoto" else local_server

    user = os.getenv("GESTION_SQL_USER") or os.getenv("SQLSERVER_USER")
    password = os.getenv("GESTION_SQL_PASSWORD") or os.getenv("SQLSERVER_PASSWORD")

    if user and password:
        return (
            f"DRIVER={{{driver}}};"
            f"SERVER={server};"
            "DATABASE=Vencorp_V2;"
            f"UID={user};"
            f"PWD={password};"
            "TrustServerCertificate=yes;"
        )

    return (
        f"DRIVER={{{driver}}};"
        f"SERVER={server};"
        "DATABASE=Vencorp_V2;"
        "Trusted_Connection=yes;"
        "TrustServerCertificate=yes;"
    )


def _normalizar_valor_sql(value: Any) -> Any:
    try:
        import pandas as pd
        import numpy as np
    except Exception:
        pd = None
        np = None

    if value is None:
        return None

    if pd is not None:
        try:
            if pd.isna(value):
                return None
        except Exception:
            pass

        try:
            if isinstance(value, pd.Timestamp):
                return value.to_pydatetime()
        except Exception:
            pass

    if np is not None:
        try:
            if isinstance(value, np.generic):
                return value.item()
        except Exception:
            pass

    return value


def _columnas_insertables_sql(cursor: Any, schema: str, table: str) -> list[str]:
    query = """
    SELECT c.name
    FROM sys.columns c
    INNER JOIN sys.objects o ON c.object_id = o.object_id
    INNER JOIN sys.schemas s ON o.schema_id = s.schema_id
    WHERE s.name = ?
      AND o.name = ?
      AND o.type = 'U'
      AND c.is_identity = 0
      AND c.is_computed = 0
    ORDER BY c.column_id;
    """

    return [row[0] for row in cursor.execute(query, schema, table).fetchall()]


def _contar_registros_tabla_sql(cursor: Any, schema: str, table: str) -> int | None:
    """
    Conteo rápido de registros usando metadatos de SQL Server.

    Evita SELECT COUNT(*) sobre tablas grandes, porque en
    gestion.gestion_adminfo_onedrive puede tardar varios minutos.

    Usa sys.dm_db_partition_stats para obtener el total de filas del heap
    o índice clustered. Si falla por permisos u otra causa, devuelve None
    para no bloquear la Fase I con un conteo completo lento.
    """
    try:
        sql = """
        SELECT SUM(row_count) AS total_rows
        FROM sys.dm_db_partition_stats
        WHERE object_id = OBJECT_ID(?)
          AND index_id IN (0, 1)
        """

        object_name = f"[{schema}].[{table}]"
        row = cursor.execute(sql, object_name).fetchone()

        if row and row[0] is not None:
            return int(row[0])

        return None

    except Exception:
        return None

# ============================================================
# CONSOLIDAR GESTIÓN FASE I - MAPEO EXACTO SQL
# ============================================================

GESTION_SQL_SCHEMA_EXACTO = [
    {
        "sql": "ClienteNro_Contrato",
        "tipo": "float",
        "aliases": ["ClienteNro_Contrato", "Cliente Nro.", "Cliente Nro", "NroCliente_Contrato", "Codigo_Cliente", "Código Cliente"],
    },
    {
        "sql": "Fecha_De_Gestion",
        "tipo": "datetime",
        "aliases": ["Fecha_De_Gestion", "Fecha De Gestion", "Fecha de Gestión", "Fecha Gestion", "Fecha"],
    },
    {
        "sql": "Hora_De_Gestion",
        "tipo": "datetime_time",
        "aliases": ["Hora_De_Gestion", "Hora De Gestion", "Hora de Gestión", "Hora Gestion", "Hora"],
    },
    {
        "sql": "Duracion_llamada",
        "tipo": "datetime_time",
        "aliases": ["Duracion_llamada", "Duracion llamada", "Duración llamada", "Duracion de llamada", "Duración de llamada"],
    },
    {
        "sql": "Descripcion_Codigo_De_Gestion",
        "tipo": "nvarchar",
        "max": 255,
        "aliases": ["Descripcion_Codigo_De_Gestion", "Descripcion Codigo De Gestion", "Descripción Código De Gestión", "Descripcion Codigo Gestion"],
    },
    {
        "sql": "Fecha_Compromiso",
        "tipo": "datetime",
        "aliases": ["Fecha_Compromiso", "Fecha Compromiso", "Fecha de Compromiso", "FechaCompromiso"],
    },
    {
        "sql": "Grabador",
        "tipo": "nvarchar",
        "max": 100,
        "aliases": ["Grabador"],
    },
    {
        "sql": "Responsable_De_Cobro",
        "tipo": "nvarchar",
        "max": 15,
        "aliases": ["Responsable_De_Cobro", "Responsable De Cobro", "Responsable Cobro"],
    },
    {
        "sql": "Telefonos",
        "tipo": "float",
        "aliases": ["Telefonos", "Teléfonos", "Telefono", "Teléfono"],
    },
    {
        "sql": "Tipo_Cartera",
        "tipo": "nvarchar",
        "max": 15,
        "aliases": ["Tipo_Cartera", "Tipo Cartera"],
    },
    {
        "sql": "Asesor",
        "tipo": "nvarchar",
        "max": 25,
        "aliases": ["Asesor"],
    },
    {
        "sql": "Antigüedad_De_La_Cartera",
        "tipo": "nvarchar",
        "max": 50,
        "aliases": ["Antigüedad_De_La_Cartera", "Antiguedad_De_La_Cartera", "Antigüedad De La Cartera", "Antiguedad De La Cartera"],
    },
    {
        "sql": "Nota_de_la_Gestion",
        "tipo": "nvarchar",
        "max": 255,
        "aliases": ["Nota_de_la_Gestion", "Nota de la Gestion", "Nota de la Gestión", "Nota Gestion", "Nota"],
    },
    {
        "sql": "Clase_de_Gestion",
        "tipo": "nvarchar",
        "max": 25,
        "aliases": ["Clase_de_Gestion", "Clase de Gestion", "Clase de Gestión"],
    },
    {
        "sql": "Causal_de_Mora_Respuesta",
        "tipo": "nvarchar",
        "max": 255,
        "aliases": ["Causal_de_Mora_Respuesta", "Causal de Mora Respuesta", "Causal de Mora", "Causal Mora Respuesta"],
    },
    {
        "sql": "Mes_Gestion",
        "tipo": "float",
        "aliases": ["Mes_Gestion", "Mes Gestion"],
    },
    {
        "sql": "origen_datos",
        "tipo": "nvarchar",
        "max": 25,
        "aliases": ["origen_datos", "Origen Datos", "origen datos"],
    },
    {
        "sql": "crm",
        "tipo": "nvarchar",
        "max": 10,
        "aliases": ["crm", "CRM"],
    },
]

def _gc_col_key(nombre: Any) -> str:
    import unicodedata

    texto = str(nombre or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    texto = re.sub(r"[^a-z0-9]+", "", texto)

    return texto


def _gc_mapa_columnas_excel(df: Any) -> dict[str, str]:
    return {
        _gc_col_key(col): col
        for col in df.columns
    }


def _gc_resolver_columna_excel(df: Any, aliases: list[str]) -> str | None:
    mapa = _gc_mapa_columnas_excel(df)

    for alias in aliases:
        key = _gc_col_key(alias)

        if key in mapa:
            return mapa[key]

    return None


def _gc_es_vacio(value: Any) -> bool:
    try:
        import pandas as pd

        if pd.isna(value):
            return True
    except Exception:
        pass

    texto = str(value if value is not None else "").strip()

    return texto == "" or texto.lower() in {"nan", "nat", "none", "null"}


def _gc_convertir_float(value: Any) -> tuple[Any, str]:
    if _gc_es_vacio(value):
        return None, ""

    texto = str(value).strip()

    # Evitar teléfonos o clientes con separadores.
    texto = texto.replace(" ", "")
    texto = texto.replace("-", "")
    texto = texto.replace("/", "")
    texto = texto.replace("(", "").replace(")", "")

    # Si viene con coma decimal.
    texto = texto.replace(",", ".")

    # Si viene como 123.0 para identificadores.
    try:
        return float(texto), ""
    except Exception:
        # Intento alternativo: extraer primer número.
        match = re.search(r"-?\d+(?:\.\d+)?", texto)

        if match:
            try:
                return float(match.group(0)), ""
            except Exception:
                pass

    return None, f"No se puede convertir a float: {value}"


def _gc_convertir_datetime(value: Any, time_only: bool = False) -> tuple[Any, str]:
    from datetime import datetime, date, time, timedelta

    if _gc_es_vacio(value):
        return None, ""

    # Si viene como time puro.
    if isinstance(value, time):
        return datetime.combine(date(1900, 1, 1), value), ""

    if isinstance(value, datetime):
        if time_only:
            return datetime.combine(date(1900, 1, 1), value.time()), ""
        return value, ""

    # Excel serial date/time.
    if isinstance(value, (int, float)):
        try:
            if time_only and 0 <= float(value) < 1:
                base = datetime(1900, 1, 1)
                return base + timedelta(days=float(value)), ""

            import pandas as pd
            dt = pd.to_datetime(value, unit="D", origin="1899-12-30", errors="coerce")

            if not pd.isna(dt):
                if time_only:
                    return datetime.combine(date(1900, 1, 1), dt.time()), ""
                return dt.to_pydatetime(), ""
        except Exception:
            pass

    texto = str(value).strip()

    # Hora o duración tipo HH:MM o HH:MM:SS.
    if time_only:
        match = re.match(r"^(\d{1,3}):(\d{2})(?::(\d{2}))?$", texto)

        if match:
            horas = int(match.group(1))
            minutos = int(match.group(2))
            segundos = int(match.group(3) or 0)

            # SQL datetime no guarda duración > 23h como hora real.
            # Para duraciones grandes se deja como fecha base + horas.
            base = datetime(1900, 1, 1)
            return base + timedelta(hours=horas, minutes=minutos, seconds=segundos), ""

    try:
        import pandas as pd

        dt = pd.to_datetime(texto, errors="coerce", dayfirst=False)

        if pd.isna(dt):
            dt = pd.to_datetime(texto, errors="coerce", dayfirst=True)

        if pd.isna(dt):
            return None, f"No se puede convertir a datetime: {value}"

        if time_only:
            return datetime.combine(date(1900, 1, 1), dt.time()), ""

        return dt.to_pydatetime(), ""
    except Exception as exc:
        return None, f"No se puede convertir a datetime: {value} | {exc}"


def _gc_convertir_texto(value: Any, max_len: int | None = None) -> tuple[Any, str]:
    if _gc_es_vacio(value):
        return None, ""

    texto = str(value).strip()

    if max_len and len(texto) > max_len:
        # Para nvarchar no conviene fallar; truncamos y lo reportamos como advertencia controlada.
        return texto[:max_len], f"Texto truncado a {max_len} caracteres"

    return texto, ""


def _gc_preparar_carga_exacta(df: Any) -> tuple[list[str], list[tuple[Any, ...]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    columnas_sql = [item["sql"] for item in GESTION_SQL_SCHEMA_EXACTO]
    rows = []
    errores = []
    advertencias = []
    mapeo = []

    for item in GESTION_SQL_SCHEMA_EXACTO:
        col_excel = _gc_resolver_columna_excel(df, item["aliases"])

        mapeo.append(
            {
                "columna_sql": item["sql"],
                "tipo_sql": item["tipo"],
                "columna_excel": col_excel or "",
                "estado": "OK" if col_excel else "SIN COLUMNA EXCEL - INSERTA NULL",
            }
        )

    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        values = []

        for item, map_item in zip(GESTION_SQL_SCHEMA_EXACTO, mapeo):
            col_sql = item["sql"]
            tipo = item["tipo"]
            col_excel = map_item["columna_excel"]

            raw = None if not col_excel else row.get(col_excel)

            if not col_excel:
                values.append(None)
                continue

            if tipo == "float":
                convertido, error = _gc_convertir_float(raw)
            elif tipo == "datetime":
                convertido, error = _gc_convertir_datetime(raw, time_only=False)
            elif tipo == "datetime_time":
                convertido, error = _gc_convertir_datetime(raw, time_only=True)
            elif tipo == "nvarchar":
                convertido, error = _gc_convertir_texto(raw, item.get("max"))
            else:
                convertido, error = raw, ""

            if error:
                registro = {
                    "fila_excel": row_idx,
                    "columna_sql": col_sql,
                    "columna_excel": col_excel,
                    "tipo_sql": tipo,
                    "valor": str(raw),
                    "mensaje": error,
                }

                if "truncado" in error.lower():
                    advertencias.append(registro)
                else:
                    errores.append(registro)

            values.append(convertido)

        rows.append(tuple(values))

    return columnas_sql, rows, errores, advertencias, mapeo

def cargar_informacion_gestion_sql(
    data_dir: str | Path,
    fecha_raw: str,
    conexion: str = "local",
) -> dict[str, Any]:
    """
    Fase I - Cargar información usando mapeo exacto a:
    Vencorp_V2.gestion.gestion_adminfo_onedrive
    """
    try:
        import pandas as pd
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pandas: {exc}",
        }

    try:
        import pyodbc
    except Exception as exc:
        return {
            "ok": False,
            "error": f"No se pudo importar pyodbc: {exc}",
        }

    from time import perf_counter
    from datetime import datetime as _dt_audit

    auditoria_servicio = []

    def _audit_start():
        return perf_counter(), _dt_audit.now().strftime("%H:%M:%S.%f")[:-3]

    def _audit_end(paso, inicio_perf, inicio_clock, detalle=""):
        fin_perf = perf_counter()
        fin_clock = _dt_audit.now().strftime("%H:%M:%S.%f")[:-3]
        auditoria_servicio.append({
            "paso": paso,
            "inicio": inicio_clock,
            "fin": fin_clock,
            "segundos": round(fin_perf - inicio_perf, 3),
            "detalle": detalle or "",
        })

    inicio_total_serv_perf, inicio_total_serv_clock = _audit_start()

    inicio_rutas_perf, inicio_rutas_clock = _audit_start()
    rutas = resolver_rutas(data_dir, fecha_raw)
    fecha = rutas["fecha"]
    conexion_normalizada = str(conexion or "local").lower()

    schema = "gestion"
    table = "gestion_adminfo_onedrive"
    database = "Vencorp_V2"

    ruta_excel = (
        Path(rutas["subcarpetas"]["05_final"])
        / f"{fecha}_Gestion_excel.xlsx"
    )

    carpeta_reportes = Path(rutas["subcarpetas"]["Reportes"])
    carpeta_reportes.mkdir(parents=True, exist_ok=True)
    ruta_reporte = carpeta_reportes / f"{fecha}_reporte_carga_sql.xlsx"

    _audit_end(
        "Resolver rutas y preparar carpetas",
        inicio_rutas_perf,
        inicio_rutas_clock,
        f"ruta_excel={ruta_excel}; ruta_reporte={ruta_reporte}",
    )

    if not ruta_excel.exists():
        return {
            "ok": False,
            "error": f"No existe archivo final. Ejecute primero Fase G: {ruta_excel}",
            "ruta_excel": str(ruta_excel),
            "conexion": conexion_normalizada,
            "database": database,
            "schema": schema,
            "table": table,
            "tabla_destino": f"{database}.{schema}.{table}",
        }

    try:
        inicio_excel_perf, inicio_excel_clock = _audit_start()

        df = pd.read_excel(ruta_excel)
        df.columns = [str(col).strip() for col in df.columns]

        filas_leidas = len(df)

        _audit_end(
            "Lectura Excel final",
            inicio_excel_perf,
            inicio_excel_clock,
            f"filas={filas_leidas}; columnas={len(df.columns)}",
        )

        if filas_leidas == 0:
            return {
                "ok": False,
                "error": "El archivo final no contiene filas para insertar.",
                "ruta_excel": str(ruta_excel),
                "conexion": conexion_normalizada,
                "tabla_destino": f"{database}.{schema}.{table}",
                "filas_leidas": 0,
            }

        inicio_prep_perf, inicio_prep_clock = _audit_start()

        columnas_sql, rows, errores_conversion, advertencias, mapeo = _gc_preparar_carga_exacta(df)

        _audit_end(
            "Preparación y conversión exacta",
            inicio_prep_perf,
            inicio_prep_clock,
            f"rows={len(rows)}; columnas_sql={len(columnas_sql)}; errores_conversion={len(errores_conversion)}; advertencias={len(advertencias)}",
        )

        # Si hay errores reales de conversión, NO se inserta.
        if errores_conversion:
            resumen = pd.DataFrame(
                [
                    {"control": "Conexión", "valor": conexion_normalizada},
                    {"control": "Base de datos", "valor": database},
                    {"control": "Tabla destino", "valor": f"{schema}.{table}"},
                    {"control": "Archivo leído", "valor": str(ruta_excel)},
                    {"control": "Filas leídas Excel", "valor": filas_leidas},
                    {"control": "Filas insertadas SQL", "valor": 0},
                    {"control": "Errores de conversión", "valor": len(errores_conversion)},
                    {"control": "Advertencias", "valor": len(advertencias)},
                    {"control": "Acción", "valor": "NO SE INSERTÓ POR ERRORES DE CONVERSIÓN"},
                ]
            )

            with pd.ExcelWriter(ruta_reporte) as writer:
                resumen.to_excel(writer, sheet_name="Resumen", index=False)
                pd.DataFrame(mapeo).to_excel(writer, sheet_name="MapeoColumnas", index=False)
                pd.DataFrame(errores_conversion).to_excel(writer, sheet_name="ErroresConversion", index=False)
                pd.DataFrame(advertencias).to_excel(writer, sheet_name="Advertencias", index=False)

            return {
                "ok": False,
                "tipo_error": "conversion",
                "fecha": fecha,
                "conexion": conexion_normalizada,
                "database": database,
                "schema": schema,
                "table": table,
                "tabla_destino": f"{database}.{schema}.{table}",
                "ruta_excel": str(ruta_excel),
                "ruta_reporte": str(ruta_reporte),
                "archivo_reporte": ruta_reporte.name,
                "filas_leidas": filas_leidas,
                "filas_insertadas": 0,
                "error": "Se encontraron valores incompatibles con la estructura SQL exacta. No se insertó ningún registro.",
                "errores_conversion": errores_conversion[:250],
                "advertencias": advertencias[:250],
                "mapeo": mapeo,
                "columnas_insertadas": columnas_sql,
                "columnas_omitidas_excel": [],
                "columnas_tabla_sin_excel": [],
                "total_antes": "",
                "total_despues": "",
            }

        conn_str = _connection_string_vencorp(conexion_normalizada)

        inicio_conn_perf, inicio_conn_clock = _audit_start()

        with pyodbc.connect(conn_str, timeout=30) as conn:
            cursor = conn.cursor()

            _audit_end(
                "Conexión SQL Server",
                inicio_conn_perf,
                inicio_conn_clock,
                f"conexion={conexion_normalizada}; database={database}",
            )

            inicio_count_antes_perf, inicio_count_antes_clock = _audit_start()

            total_antes = _contar_registros_tabla_sql(cursor, schema, table)

            _audit_end(
                "Conteo total antes",
                inicio_count_antes_perf,
                inicio_count_antes_clock,
                f"total_antes={total_antes}",
            )

            inicio_dup_perf, inicio_dup_clock = _audit_start()

            control_fecha = _gc_existe_fecha_gestion_sql(cursor, schema, table, fecha)

            _audit_end(
                "Control anti-duplicado por Fecha_De_Gestion",
                inicio_dup_perf,
                inicio_dup_clock,
                f"existe={control_fecha.get('existe')}; rango={control_fecha.get('fecha_inicio')} a {control_fecha.get('fecha_fin')}",
            )

            if control_fecha.get("existe"):
                _audit_end(
                    "Total servicio cargar_informacion_gestion_sql",
                    inicio_total_serv_perf,
                    inicio_total_serv_clock,
                    "Servicio finalizado sin insertar por control anti-duplicado.",
                )

                return {
                    "ok": False,
                    "tipo_error": "duplicado_fecha",
                    "fecha": fecha,
                    "conexion": conexion_normalizada,
                    "database": database,
                    "schema": schema,
                    "table": table,
                    "tabla_destino": f"{database}.{schema}.{table}",
                    "ruta_excel": str(ruta_excel),
                    "ruta_reporte": str(ruta_reporte),
                    "archivo_excel": ruta_excel.name,
                    "filas_leidas": filas_leidas,
                    "filas_insertadas": 0,
                    "columnas_insertadas": columnas_sql,
                    "columnas_omitidas_excel": [],
                    "columnas_tabla_sin_excel": [],
                    "advertencias": advertencias[:250],
                    "mapeo": mapeo,
                    "total_antes": total_antes,
                    "total_despues": total_antes,
                    "errores": [],
                    "control_duplicado_fecha": control_fecha,
                    "error": control_fecha.get("mensaje"),
                    "auditoria_servicio": auditoria_servicio,
                }

            placeholders = ", ".join(["?"] * len(columnas_sql))
            columnas_sql_brackets = ", ".join(f"[{col}]" for col in columnas_sql)

            insert_sql = (
                f"INSERT INTO [{schema}].[{table}] "
                f"({columnas_sql_brackets}) VALUES ({placeholders})"
            )

            cursor.fast_executemany = True

            filas_insertadas = 0
            errores_sql = []
            chunk_size = 1000

            inicio_insert_perf, inicio_insert_clock = _audit_start()

            for inicio in range(0, len(rows), chunk_size):
                chunk = rows[inicio: inicio + chunk_size]

                try:
                    cursor.executemany(insert_sql, chunk)
                    filas_insertadas += len(chunk)
                except Exception as exc:
                    errores_sql.append(
                        {
                            "bloque_inicio": inicio + 1,
                            "bloque_fin": inicio + len(chunk),
                            "error": str(exc),
                        }
                    )
                    conn.rollback()
                    raise

            _audit_end(
                "Inserción SQL por bloques",
                inicio_insert_perf,
                inicio_insert_clock,
                f"filas_insertadas={filas_insertadas}; chunk_size={chunk_size}",
            )

            inicio_commit_perf, inicio_commit_clock = _audit_start()

            conn.commit()

            _audit_end(
                "Commit SQL",
                inicio_commit_perf,
                inicio_commit_clock,
                "Confirmación de transacción SQL.",
            )

            inicio_count_despues_perf, inicio_count_despues_clock = _audit_start()

            total_despues = _contar_registros_tabla_sql(cursor, schema, table)

            _audit_end(
                "Conteo total después",
                inicio_count_despues_perf,
                inicio_count_despues_clock,
                f"total_despues={total_despues}",
            )

        resumen = pd.DataFrame(
            [
                {"control": "Conexión", "valor": conexion_normalizada},
                {"control": "Base de datos", "valor": database},
                {"control": "Tabla destino", "valor": f"{schema}.{table}"},
                {"control": "Archivo leído", "valor": str(ruta_excel)},
                {"control": "Filas leídas Excel", "valor": filas_leidas},
                {"control": "Filas insertadas SQL", "valor": filas_insertadas},
                {"control": "Columnas insertadas", "valor": len(columnas_sql)},
                {"control": "Advertencias", "valor": len(advertencias)},
                {"control": "Total tabla antes", "valor": "" if total_antes is None else total_antes},
                {"control": "Total tabla después", "valor": "" if total_despues is None else total_despues},
                {"control": "Errores SQL", "valor": len(errores_sql)},
            ]
        )

        inicio_reporte_perf, inicio_reporte_clock = _audit_start()

        with pd.ExcelWriter(ruta_reporte) as writer:
            resumen.to_excel(writer, sheet_name="Resumen", index=False)
            pd.DataFrame(mapeo).to_excel(writer, sheet_name="MapeoColumnas", index=False)
            pd.DataFrame({"columna_insertada": columnas_sql}).to_excel(writer, sheet_name="ColumnasInsertadas", index=False)
            pd.DataFrame(advertencias).to_excel(writer, sheet_name="Advertencias", index=False)
            pd.DataFrame(errores_sql).to_excel(writer, sheet_name="ErroresSQL", index=False)

        _audit_end(
            "Generación reporte Excel",
            inicio_reporte_perf,
            inicio_reporte_clock,
            f"ruta_reporte={ruta_reporte}",
        )

        _audit_end(
            "Total servicio cargar_informacion_gestion_sql",
            inicio_total_serv_perf,
            inicio_total_serv_clock,
            "Tiempo total dentro del servicio.",
        )

        return {
            "ok": True,
            "fecha": fecha,
            "conexion": conexion_normalizada,
            "database": database,
            "schema": schema,
            "table": table,
            "tabla_destino": f"{database}.{schema}.{table}",
            "ruta_excel": str(ruta_excel),
            "ruta_reporte": str(ruta_reporte),
            "archivo_excel": ruta_excel.name,
            "archivo_reporte": ruta_reporte.name,
            "filas_leidas": filas_leidas,
            "filas_insertadas": filas_insertadas,
            "columnas_insertadas": columnas_sql,
            "columnas_omitidas_excel": [],
            "columnas_tabla_sin_excel": [],
            "advertencias": advertencias[:250],
            "mapeo": mapeo,
            "total_antes": total_antes,
            "total_despues": total_despues,
            "errores": errores_sql,
            "auditoria_servicio": auditoria_servicio,
        }
    except Exception as exc:
        return {
            "ok": False,
            "tipo_error": "sql",
            "fecha": fecha,
            "conexion": conexion_normalizada,
            "database": database,
            "schema": schema,
            "table": table,
            "tabla_destino": f"{database}.{schema}.{table}",
            "ruta_excel": str(ruta_excel),
            "ruta_reporte": str(ruta_reporte),
            "error": str(exc),
        }

