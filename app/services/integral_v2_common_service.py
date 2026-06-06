# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import os
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path
from time import perf_counter
from typing import Any

from app.services.integral_connection_service import (
    IntegralConnectionFactory,
    construir_contexto_integral_v2,
)


def now_ms() -> str:
    return datetime.now().strftime("%H:%M:%S.%f")[:-3]


def parse_fecha(value: str) -> datetime:
    raw = str(value or "").strip()
    for fmt in ("%Y%m%d", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    raise ValueError(f"Fecha de proceso no válida: {value}")


def fecha_tokens(value: str) -> dict[str, str]:
    dt = parse_fecha(value)
    return {
        "yyyymmdd": dt.strftime("%Y%m%d"),
        "ddmmyyyy": dt.strftime("%d%m%Y"),
        "dd-mm-yyyy": dt.strftime("%d-%m-%Y"),
        "yyyy-mm-dd": dt.strftime("%Y-%m-%d"),
        "sql": dt.strftime("%Y-%m-%d"),
        "visual": dt.strftime("%d/%m/%Y"),
    }


def data_dir() -> Path:
    return Path(os.getenv("INTEGRAL_DATA_DIR") or os.getenv("DATA_DIR") or r"E:\data")


def ruta_busqueda(conexion: str) -> Path:
    if str(conexion or "local").lower() == "remoto":
        return Path(
            os.getenv("INTEGRAL_RUTA_BUSQUEDA_REMOTO")
            or r"\\10.24.90.118\Vencorp\COBRANZA %\2024\ Prueba_carga_diaria_Aster_voip\Cobranzas Integrales"
        )
    return Path(
        os.getenv("INTEGRAL_RUTA_BUSQUEDA_LOCAL")
        or r"D:\Develop\ETL\Nicaragua_Proceso\unidad_red_integral\Cobranzas Integrales"
    )


def trabajo_dir(fecha_yyyymmdd: str) -> Path:
    return data_dir() / fecha_yyyymmdd / "Integral"


def original_dir(fecha_yyyymmdd: str) -> Path:
    path = trabajo_dir(fecha_yyyymmdd) / "01_original"
    path.mkdir(parents=True, exist_ok=True)
    return path


def normalize_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text)


def find_column(columns: list[str], variants: list[str]) -> str | None:
    lookup = {normalize_text(c): c for c in columns}
    for variant in variants:
        key = normalize_text(variant)
        if key in lookup:
            return lookup[key]
    return None


def unique_columns(raw_columns: list[Any]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, col in enumerate(raw_columns, start=1):
        base = str(col or "").strip() or f"columna_{i}"
        key = normalize_text(base)
        count = seen.get(key, 0)
        name = base if count == 0 else f"{base}_{count + 1}"
        seen[key] = count + 1
        out.append(name)
    return out


def read_table(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    ext = path.suffix.lower()
    if ext in {".xlsx", ".xlsm"}:
        return read_excel(path)
    if ext in {".csv", ".txt"}:
        return read_csv_file(path)
    raise ValueError(f"Extensión no soportada: {ext}")


def read_excel(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header: list[str] | None = None
    header_row = 1

    for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [str(v).strip() if v is not None else "" for v in row]
        if len([v for v in values if v]) >= 2:
            header = unique_columns(values)
            header_row = row_num
            break
        if row_num >= 25:
            break

    if not header:
        wb.close()
        raise ValueError("No se detectó encabezado en el archivo.")

    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        item: dict[str, Any] = {}
        has_data = False
        for idx, col in enumerate(header):
            value = row[idx] if idx < len(row) else None
            if value not in (None, ""):
                has_data = True
            item[col] = value
        if has_data:
            rows.append(item)

    wb.close()
    return header, rows


def read_csv_file(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "latin-1", "cp1252"):
        try:
            with path.open("r", encoding=encoding, newline="") as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except Exception:
                    dialect = csv.excel
                reader = csv.DictReader(fh, dialect=dialect)
                if not reader.fieldnames:
                    raise ValueError("No se detectó encabezado en CSV/TXT.")
                columns = unique_columns(reader.fieldnames)
                rows: list[dict[str, Any]] = []
                for row in reader:
                    item: dict[str, Any] = {}
                    has_data = False
                    for original, col in zip(reader.fieldnames, columns):
                        value = row.get(original)
                        if value not in (None, ""):
                            has_data = True
                        item[col] = value
                    if has_data:
                        rows.append(item)
                return columns, rows
        except Exception as exc:
            last_error = exc
    raise last_error or ValueError("No se pudo leer CSV/TXT.")


def latest_original_file(fecha_yyyymmdd: str) -> Path:
    folder = original_dir(fecha_yyyymmdd)
    candidates = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in {".xlsx", ".xlsm", ".csv", ".txt"}]
    if not candidates:
        raise FileNotFoundError(f"No hay archivo copiado en {folder}")
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0]


def resolve_file_from_form(form, fecha_yyyymmdd: str) -> Path:
    for key in ("archivo", "archivo_candidato", "archivo_seleccionado", "integral_archivo"):
        raw = str(form.get(key) or "").strip()
        if raw:
            path = Path(raw)
            if path.exists():
                return path
    return latest_original_file(fecha_yyyymmdd)


def step(codigo: str, paso: str, entrada: Any, func):
    inicio = now_ms()
    t0 = perf_counter()
    try:
        salida, mensaje = func()
        estado = "ok"
        advertencia = None
    except Exception as exc:
        salida = "-"
        mensaje = str(exc)
        estado = "error"
        advertencia = f"{paso}: {mensaje}"
    fin = now_ms()
    return {
        "#": None,
        "Código": codigo,
        "Paso": paso,
        "Estado": estado,
        "Inicio": inicio,
        "Fin": fin,
        "Seg.": f"{perf_counter() - t0:.2f}",
        "Entrada": entrada,
        "Salida": salida,
        "Mensaje": mensaje,
    }, advertencia


def enumerate_steps(tiempos: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for i, item in enumerate(tiempos, start=1):
        item["#"] = i
    return tiempos


def source_sql_query(conexion: str, fecha_sql: str) -> str:
    engine = IntegralConnectionFactory.source_engine(conexion)
    if engine == "mysql":
        return (
            "SELECT entidad, COUNT(1) AS registros "
            "FROM comentarios "
            f"WHERE fecha >= '{fecha_sql}' "
            f"AND fecha < DATE_ADD('{fecha_sql}', INTERVAL 1 DAY) "
            "AND (entidad LIKE '%reven%' OR entidad LIKE '%anual%') "
            "GROUP BY entidad ORDER BY entidad"
        )
    return (
        f"DECLARE @fecha_sql date = '{fecha_sql}'; "
        "SELECT entidad, COUNT_BIG(1) AS registros "
        "FROM dbo.comentarios "
        "WHERE fecha >= @fecha_sql "
        "AND fecha < DATEADD(day, 1, @fecha_sql) "
        "AND (entidad LIKE '%reven%' OR entidad LIKE '%anual%') "
        "GROUP BY entidad ORDER BY entidad;"
    )


def fetch_sql_entities(conexion: str, fecha_sql: str) -> list[dict[str, Any]]:
    conn = IntegralConnectionFactory.connect_source_comentarios(conexion)
    try:
        cur = conn.cursor()
        cur.execute(source_sql_query(conexion, fecha_sql))
        rows = cur.fetchall()
        result = []
        for row in rows:
            entidad = str(row[0] or "").strip()
            registros = int(row[1] or 0)
            if entidad:
                result.append({"entidad": entidad, "registros": registros})
        return result
    finally:
        try:
            conn.close()
        except Exception:
            pass


def count_file_entities(file_path: Path) -> tuple[list[dict[str, Any]], int, str]:
    columns, rows = read_table(file_path)
    entity_col = find_column(columns, ["Entidad", "entidad", "ENTIDAD"])
    if not entity_col:
        raise ValueError("El archivo no contiene columna Entidad.")
    counts: dict[str, int] = {}
    for row in rows:
        entidad = str(row.get(entity_col) or "").strip()
        if entidad:
            counts[entidad] = counts.get(entidad, 0) + 1
    data = [{"entidad": k, "registros": v} for k, v in counts.items()]
    data.sort(key=lambda x: x["entidad"].lower())
    return data, len(rows), entity_col


def compare_entities(sql_data: list[dict[str, Any]], file_data: list[dict[str, Any]]) -> list[dict[str, Any]]:
    sql_map = {x["entidad"]: int(x["registros"]) for x in sql_data}
    file_map = {x["entidad"]: int(x["registros"]) for x in file_data}
    names = sorted(set(sql_map) | set(file_map), key=lambda x: x.lower())
    rows = []
    for name in names:
        reg_sql = sql_map.get(name, 0)
        reg_file = file_map.get(name, 0)
        diff = reg_sql - reg_file
        if reg_sql and reg_file and diff == 0:
            estado = "OK"
            seleccionar = True
        elif reg_sql and reg_file:
            estado = "DIFERENCIA"
            seleccionar = False
        elif reg_sql and not reg_file:
            estado = "SOLO SQL"
            seleccionar = False
        else:
            estado = "SOLO ARCHIVO"
            seleccionar = False
        rows.append({
            "valor_entidad": name,
            "entidad_sql": name if reg_sql else "-",
            "registros_sql": reg_sql,
            "entidad_archivo": name if reg_file else "-",
            "registros_archivo": reg_file,
            "diferencia": diff,
            "estado": estado,
            "seleccionar": seleccionar,
        })
    return rows
